import openpyxl, json, re, sys, collections, os
src, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src)
ws = wb['Master_Combined_Mocks']
hdr = [c.value for c in ws[1]]
rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
def s(v):
    if v is None: return ""
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return str(v).strip()
AI = re.compile(r'clou?de|colude|gpt', re.I)
def tag(srcv):
    t = s(srcv).replace("\n", " ")
    t = re.sub(r'\s+', ' ', t)
    return "" if (not t or AI.search(t)) else t
qs = []
for r in rows:
    ans = int(float(r[13]))
    qs.append(dict(id=s(r[0]), q=s(r[1]), o=[s(r[2]), s(r[3]), s(r[4]), s(r[5])], a=ans,
                   e=s(r[7]), topic=s(r[9]), concept=s(r[10]), src=tag(r[11]),
                   lvl=re.sub(r'^\d+\.\s*', '', s(r[14])), mock=s(r[16]), qno=int(float(r[17] or 0))))
# topics ordered by first appearance of Source_Topic_No
order = collections.OrderedDict()
for r in sorted(rows, key=lambda r: float(r[15] or 999)):
    order.setdefault(s(r[9]), None)
topics = []
for i, name in enumerate(order, 1):
    tid = f"hi-{i:02d}"
    items = [{k: v for k, v in q.items() if k not in ('topic', 'mock', 'qno')} for q in qs if q['topic'] == name]
    json.dump({"id": tid, "name": name, "subject": "hindi", "board": "RPSC", "questions": items},
              open(f"{out}/data/topics/{tid}.json", "w"), ensure_ascii=False, separators=(',', ':'))
    topics.append({"id": tid, "name": name, "count": len(items)})
mocks = []
for n in range(1, 11):
    mid = f"SI_Mock_{n}"
    sh = wb[f'SI_Mock-{n}']
    m = [c.value for c in sh[2]]
    items = sorted([q for q in qs if q['mock'] == mid], key=lambda q: q['qno'])
    meta = dict(id=f"si-mock-{n:02d}", name=f"SI 2021 (Re-exam) Mock {n}", paper="हिंदी",
                header=s(m[27]).strip("✍️ ") or f"SI_Mock Test-{n}",
                minutes=int(float(m[20] or 120)), pos=float(m[21] or 2), neg=float(m[22] or 0.66),
                optionE=s(m[23]) or "अनुत्तरित प्रश्न", insHi=s(m[24]), insEn=s(m[25]), count=len(items))
    json.dump(dict(meta, questions=[{k: v for k, v in q.items() if k not in ('mock', 'qno')} for q in items]),
              open(f"{out}/data/mocks/{meta['id']}.json", "w"), ensure_ascii=False, separators=(',', ':'))
    mocks.append(meta | {"insHi": None, "insEn": None})
    for k in ("insHi", "insEn"): mocks[-1].pop(k)
soon = [("rj-geo", "राजस्थान भूगोल", "🏜️"), ("rj-his", "राजस्थान इतिहास", "🏰"), ("rj-art", "कला एवं संस्कृति", "🎨"),
        ("polity", "भारतीय राजव्यवस्था", "⚖️"), ("science", "सामान्य विज्ञान", "🔬"), ("reasoning", "तर्कशक्ति एवं गणित", "🧮")]
catalog = {"updated": "2026-10-10",
    "subjects": [{"id": "hindi", "name": "हिंदी", "icon": "अ", "boards": {"RPSC": topics, "RSSB": []}}] +
                [{"id": i, "name": n, "icon": ic, "boards": {"RPSC": [], "RSSB": []}} for i, n, ic in soon],
    "mocks": mocks}
json.dump(catalog, open(f"{out}/data/catalog.json", "w"), ensure_ascii=False, indent=1)
print(len(qs), "questions,", len(topics), "topics,", len(mocks), "mocks")
